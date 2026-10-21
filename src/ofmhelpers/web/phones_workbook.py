"""
The /phones Excel export: everything a VA needs to sign in on a phone, in
plaintext on purpose (decided 2026-10-20; the file is for the VAs).

Two sheets:
- **Phones** -- one row per phone: its details, then per service any
  phone uses (in services.json order) the usernames and passwords, plus
  backup codes for social services. Several logins of one service share a
  cell, one per line; a custom service's login reads "Provider: username".
- **Logins** -- one row per login per phone it is on, with that phone's
  details: the sheet to filter and sort.

Every cell is written as a string: openpyxl would otherwise store a value
starting with "=" (a password can) as a formula Excel evaluates.
"""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ofmhelpers.phones import ServiceKind, services

if TYPE_CHECKING:
    from openpyxl.worksheet._write_only import WriteOnlyWorksheet

    from ofmhelpers.web.schemas import Device, DeviceLogin

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_PHONE_HEADERS = [
    "Phone",
    "Phone number",
    "Model",
    "Color",
    "Owner",
]
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _phone_cells(device: Device) -> list[str]:
    return [
        device.label,
        device.phone_number,
        device.model_label,
        device.finish.label if device.finish else "",
        device.owner,
    ]


def _name(login: DeviceLogin) -> str:
    if login.provider:
        return f"{login.provider}: {login.username}"
    return login.username


def _row(sheet: WriteOnlyWorksheet, values: list[str], header: bool = False) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        if header:
            cell.font = _HEADER_FONT
        else:
            cell.alignment = _WRAP
        cells.append(cell)
    sheet.append(cells)


def _widths(sheet: WriteOnlyWorksheet, count: int, width: int = 22) -> None:
    """Write-only sheets take column widths before the first row."""
    for index in range(1, count + 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _phones_sheet(book: Workbook, devices: list[Device]) -> None:
    used = {x.service for d in devices for x in d.logins}
    listed = [s for key, s in services().items() if key in used]
    headers = [*_PHONE_HEADERS]
    for service in listed:
        headers += [service.label, f"{service.label} password"]
        if service.kind is ServiceKind.SOCIAL:
            headers.append(f"{service.label} backup codes")
    headers.append("Notes")
    sheet = book.create_sheet("Phones")
    _widths(sheet, len(headers))
    sheet.freeze_panes = "B2"
    _row(sheet, headers, header=True)
    for device in devices:
        values = _phone_cells(device)
        for service in listed:
            logins = [x for x in device.logins if x.service == service.key]
            values.append("\n".join(_name(login) for login in logins))
            values.append("\n".join(login.password for login in logins))
            if service.kind is ServiceKind.SOCIAL:
                values.append("\n\n".join(login.backup_codes for login in logins))
        values.append(device.notes)
        _row(sheet, values)


def _logins_sheet(book: Workbook, devices: list[Device]) -> None:
    headers = [
        *_PHONE_HEADERS,
        "Service",
        "Username",
        "Password",
        "Backup codes",
        "Notes",
    ]
    sheet = book.create_sheet("Logins")
    _widths(sheet, len(headers))
    sheet.freeze_panes = "A2"
    _row(sheet, headers, header=True)
    for device in devices:
        for login in device.logins:
            _row(
                sheet,
                [
                    *_phone_cells(device),
                    login.label,
                    login.username,
                    login.password,
                    login.backup_codes,
                    login.notes,
                ],
            )


def build(devices: list[Device]) -> bytes:
    """The .xlsx file of `devices`, in the order given."""
    book = Workbook(write_only=True)
    _phones_sheet(book, devices)
    _logins_sheet(book, devices)
    out = BytesIO()
    book.save(out)
    return out.getvalue()
