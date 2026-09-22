"""Writes qualified leads to a formatted .xlsx for manual outreach.

Its own exporter rather than a widened PostExcelExporter: that one's header,
row shape and link column are the post grid the scraper/ranker pipeline
feeds it, and a lead sheet shares none of those columns. Two small writers
beat one that branches on which caller it has.

Same fallback contract as PostExcelExporter: a sheet that cannot be saved
is dumped to .csv next to it rather than lost, because the run that
produced it cost Apify credits.
"""

from __future__ import annotations

import contextlib
import csv
from pathlib import Path
from typing import TYPE_CHECKING, ClassVar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ofmhelpers.log import get_logger

logger = get_logger(__name__)

if TYPE_CHECKING:
    from ofmhelpers.scraping.models import Lead


class LeadExcelExporter:
    HEADER: ClassVar[list[str]] = [
        "Username",
        "Profile",
        "Followers",
        "Confidence",
        "OnlyFans",
        "Evidence",
        "Name",
        "Private",
        "Bio",
    ]
    COL_WIDTHS: ClassVar[list[int]] = [20, 40, 12, 12, 45, 38, 22, 9, 70]

    # 1-based indexes into HEADER of the two columns rendered as links.
    PROFILE_COL = 2
    ONLYFANS_COL = 5
    # Bio and evidence are the only columns long enough to need wrapping.
    WRAP_COLS = (6, 9)

    HEADER_FILL = PatternFill("solid", start_color="1a1a2e")
    HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    ALT_FILL = PatternFill("solid", start_color="F2F2F7")
    BODY_FONT = Font(name="Arial", size=10)
    LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")

    def _lead_row(self, lead: Lead) -> list:
        profile = lead.profile
        return [
            profile.username,
            profile.profile_url,
            profile.followers,
            lead.confidence.value,
            lead.onlyfans_url or "",
            lead.evidence,
            profile.full_name or "",
            "yes" if profile.is_private else "",
            (profile.biography or "").replace("\n", " "),
        ]

    def export(self, leads: list[Lead], output_path: str) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "leads"
        self._write_sheet(worksheet, leads)

        try:
            workbook.save(output_path)
            logger.info("saved %d leads -> %s", len(leads), output_path)
        except Exception:
            csv_path = self._fallback_csv(leads, output_path)
            logger.exception("failed to save xlsx -- dumped to %s", csv_path)

    def _write_sheet(self, worksheet, leads: list[Lead]) -> None:
        worksheet.append(self.HEADER)

        for col_idx, (cell, width) in enumerate(
            zip(worksheet[1], self.COL_WIDTHS, strict=False), start=1
        ):
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            worksheet.column_dimensions[get_column_letter(col_idx)].width = width
        worksheet.row_dimensions[1].height = 22

        for row_idx, lead in enumerate(leads, start=2):
            try:
                worksheet.append(self._lead_row(lead))
            except Exception:
                logger.warning(
                    "skipped row %d (%s)", row_idx, lead.profile.username, exc_info=True
                )
                continue

            fill = self.ALT_FILL if row_idx % 2 == 0 else None
            link_cols = (self.PROFILE_COL, self.ONLYFANS_COL)
            for col_idx, cell in enumerate(worksheet[row_idx], start=1):
                cell.font = self.LINK_FONT if col_idx in link_cols else self.BODY_FONT
                cell.alignment = Alignment(
                    vertical="center", wrap_text=col_idx in self.WRAP_COLS
                )
                if fill:
                    cell.fill = fill

            for col_idx, url in (
                (self.PROFILE_COL, lead.profile.profile_url),
                (self.ONLYFANS_COL, lead.onlyfans_url),
            ):
                if url and url.startswith("http"):
                    with contextlib.suppress(Exception):
                        worksheet.cell(row=row_idx, column=col_idx).hyperlink = url

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    def _fallback_csv(self, leads: list[Lead], xlsx_path: str) -> str:
        """Dumps the same rows next to the .xlsx that could not be written and
        returns where they went. The caller logs it -- it is the one holding
        the exception that made this necessary."""
        csv_path = xlsx_path.replace(".xlsx", ".csv")
        with Path(csv_path).open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(self.HEADER)
            writer.writerows(self._lead_row(lead) for lead in leads)
        return csv_path
