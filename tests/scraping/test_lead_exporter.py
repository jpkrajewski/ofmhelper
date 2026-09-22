"""The lead sheet: the artifact an Apify run was paid for, so the contract
is that rows survive -- including when the .xlsx itself cannot be written."""

from pathlib import Path

from openpyxl import load_workbook

from ofmhelpers.scraping.lead_exporter import LeadExcelExporter
from ofmhelpers.scraping.models import Confidence, InstagramProfile, Lead


def make_lead(username="somegirl", confidence=Confidence.CERTAIN, of_url=None) -> Lead:
    return Lead(
        profile=InstagramProfile(
            username=username,
            full_name="Some Girl",
            followers=4200,
            biography="line one\nline two",
        ),
        confidence=confidence,
        onlyfans_url=of_url,
        evidence="onlyfans link in bio",
    )


def test_export_writes_a_row_per_lead(tmp_path):
    out = tmp_path / "leads.xlsx"
    LeadExcelExporter().export(
        [
            make_lead(of_url="https://onlyfans.com/somegirl"),
            make_lead(username="other", confidence=Confidence.LIKELY),
        ],
        str(out),
    )

    sheet = load_workbook(out).active
    assert sheet.max_row == 3  # header + two leads
    assert [cell.value for cell in sheet[1]] == LeadExcelExporter.HEADER

    first = {
        header: cell.value
        for header, cell in zip(LeadExcelExporter.HEADER, sheet[2], strict=True)
    }
    assert first["Username"] == "somegirl"
    assert first["Profile"] == "https://www.instagram.com/somegirl/"
    assert first["Followers"] == 4200
    assert first["Confidence"] == "CERTAIN"
    assert first["OnlyFans"] == "https://onlyfans.com/somegirl"
    # Newlines in a bio would break the csv fallback's row alignment.
    assert "\n" not in first["Bio"]


def test_export_hyperlinks_both_link_columns(tmp_path):
    out = tmp_path / "leads.xlsx"
    LeadExcelExporter().export(
        [make_lead(of_url="https://onlyfans.com/somegirl")], str(out)
    )

    sheet = load_workbook(out).active
    assert sheet.cell(row=2, column=LeadExcelExporter.PROFILE_COL).hyperlink is not None
    assert (
        sheet.cell(row=2, column=LeadExcelExporter.ONLYFANS_COL).hyperlink is not None
    )


def test_export_leaves_the_onlyfans_cell_unlinked_on_a_likely_lead(tmp_path):
    out = tmp_path / "leads.xlsx"
    LeadExcelExporter().export([make_lead(confidence=Confidence.LIKELY)], str(out))

    sheet = load_workbook(out).active
    assert sheet.cell(row=2, column=LeadExcelExporter.ONLYFANS_COL).hyperlink is None


def test_export_falls_back_to_csv_when_the_xlsx_cannot_be_saved(tmp_path, monkeypatch):
    out = tmp_path / "leads.xlsx"

    def explode(*_args, **_kwargs):
        msg = "disk full"
        raise OSError(msg)

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", explode)
    LeadExcelExporter().export([make_lead()], str(out))

    csv_path = tmp_path / "leads.csv"
    assert csv_path.is_file()
    assert "somegirl" in csv_path.read_text(encoding="utf-8")


def test_export_of_nothing_still_writes_a_header(tmp_path):
    out = tmp_path / "leads.xlsx"
    LeadExcelExporter().export([], str(out))

    sheet = load_workbook(out).active
    assert sheet.max_row == 1
    assert Path(out).is_file()
